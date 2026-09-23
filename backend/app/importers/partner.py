"""Versioned reader for the original HackAlem 22.09.2026 partner exports.

This is a file exchange profile, not a generic 1C connector. Original files are
read without alteration. Unknown quantities and units are never filled with zero.
"""
from collections import defaultdict
from datetime import date, datetime
from hashlib import sha256
import csv
import re

from openpyxl import load_workbook

from ..engine.models import (Inventory, MonthlySale, ProductInput, Sale, Seasonality,
                             Shipment, Source, SupplierPolicy)
from ..engine.service import make_dataset
from .xlsx import number as parse_number

PROFILE = 'hackalem-2026-v1'
SNAPSHOT = date(2026, 9, 22)
SCOPE = 'Все склады'
MONTHS = ['янв', 'фев', 'мар', 'апр', 'май', 'июн', 'июл', 'авг', 'сен', 'окт', 'ноя', 'дек']


def number(value):
    if isinstance(value, str) and value.startswith(('#N/A', '#VALUE!', '#REF!', '#DIV/0!', '#NAME?', '#NUM!', '#NULL!')):
        return None
    return parse_number(value)


class PartnerFormatError(ValueError):
    """Only controlled, user-facing format messages may leave the importer."""


def kind_of(name):
    name = name.casefold()
    for marker, kind in [('moq', 'moq'), ('динамика', 'sales'), ('ежемесячные продажи', 'monthly_sales'),
                         ('ежемесячные остатки', 'monthly_stock'), ('сезонность', 'seasonality'),
                         ('путь', 'incoming'), ('в пути', 'incoming')]:
        if marker in name:
            return kind
    if name == 'inventory.csv':
        return 'inventory'
    raise PartnerFormatError('Неизвестный файл. Выберите шесть исходных выгрузок одного поставщика; дополнительный остаток — inventory.csv.')


def month_columns(header):
    result = []
    for index, value in enumerate(header):
        match = re.fullmatch(r'([а-яё]+)\.? (20\d{2})', str(value).strip().lower())
        if match and match[1][:3] in MONTHS:
            result.append((index, date(int(match[2]), MONTHS.index(match[1][:3]) + 1, 1)))
    if not result or len({m for _, m in result}) != len(result):
        raise PartnerFormatError('Не распознаны уникальные месячные столбцы отчёта.')
    return result


def import_partner(request):
    from .service import _check_archive, canonical_supplier
    supplier = canonical_supplier(request.supplier_id)
    if request.as_of != SNAPSHOT or request.warehouse_id != SCOPE:
        raise PartnerFormatError('Профиль исходных файлов: дата 22.09.2026, охват «Все склады». Для других дат используйте явный XLSX-профиль.')
    files = {}
    for f in request.files:
        if sha256(f.path.read_bytes()).hexdigest() != f.sha256:
            raise PartnerFormatError('Контрольная сумма файла не совпала.')
        kind = kind_of(f.original_name)
        if kind in files:
            raise PartnerFormatError('Загружены два файла одного типа. Выберите один комплект поставщика.')
        files[kind] = f
    required = {'sales', 'monthly_sales', 'monthly_stock', 'incoming', 'seasonality', 'moq'}
    if required - files.keys():
        raise PartnerFormatError('Нужны все шесть файлов: динамика, продажи по месяцам, остатки по месяцам, товары в пути, сезонность и MOQ.')
    products = defaultdict(lambda: {'units': set(), 'sales': [], 'monthly': {}, 'stocks': [], 'incoming': [], 'warnings': set()})
    audit, seasonality = [], None

    def item(sku):
        # Exported 1C identifiers must be text, including leading zeroes and underscores.
        if not isinstance(sku, str) or not sku.strip() or sku.strip().lower() in ('итого', 'всего'):
            raise PartnerFormatError('Код 1С должен быть текстом. Числовой код или строка итога не могут использоваться как товар.')
        return products[sku]

    for kind in ('moq', 'monthly_stock', 'monthly_sales', 'incoming', 'seasonality', 'sales'):
        f = files[kind]
        if f.path.suffix.lower() != '.xlsx':
            raise PartnerFormatError('Шесть исходных отчётов должны быть в формате XLSX.')
        _check_archive(f.path)
        book = load_workbook(f.path, read_only=True, data_only=True, keep_links=False)
        try:
            sheet = book['TDSheet'] if kind == 'incoming' and supplier == 'systeme_electric' and 'TDSheet' in book.sheetnames else book.worksheets[0]
            rows = iter(sheet.iter_rows(values_only=True))
            header = tuple(str(v).strip() if v is not None else '' for v in next(rows))
            header_row = 1
            if kind == 'incoming' and supplier == 'systeme_electric':
                header = tuple(str(v).strip() if v is not None else '' for v in next(rows))
                header_row = 2
            def col(label):
                if header.count(label) != 1:
                    raise PartnerFormatError(f'Не найден однозначный столбец «{label}» в отчёте {kind}. Выберите исходный файл нужного поставщика.')
                return header.index(label)
            def origin(n):
                return f'{f.original_name}:{sheet.title}:{n}'
            seen = used = 0
            rejected = []
            if kind == 'seasonality':
                values = list(rows)
                if values[8][11] != 'СЕЗОННОСТЬ':
                    raise PartnerFormatError('Ожидался итоговый столбец «СЕЗОННОСТЬ» L10.')
                coefficients = tuple(number(values[m + 9][11]) for m in range(12))
                if any(v is None or v <= 0 for v in coefficients):
                    raise PartnerFormatError('Нужны 12 положительных рассчитанных коэффициентов сезонности. Пересохраните файл с вычисленными формулами.')
                seasonality = Seasonality(coefficients=coefficients, source=Source(origin=origin(11), status='estimated'))
                seen = used = 12
            else:
                sku_col = col('Код') if kind == 'sales' else col('Код 1с') if kind == 'incoming' or (kind == 'moq' and supplier == 'iek') else col('Номенклатура.Код')
                wide = month_columns(header) if kind in ('monthly_stock', 'monthly_sales') else []
                if kind == 'sales':
                    for label in ('Дата', 'Номер', 'Документ', 'Номенклатура', 'Ед.', 'Склад', 'Количество'): col(label)
                if kind == 'moq':
                    name_col = col('Наименование' if supplier == 'iek' else 'Номенклатура')
                    article_col = col('Артикул поставщика' if supplier == 'iek' else 'Артикул')
                    rule_col = col('Мин. разр. к отгр.' if supplier == 'iek' else 'Кратность')
                if kind == 'incoming':
                    incoming_columns = []
                    if supplier == 'iek':
                        col('Артикул ИЭК')
                        for i, label in enumerate(header[3:], 3):
                            match = re.search(r'поступление до (\d{2}\.\d{2}\.\d{4})', label)
                            if not match:
                                raise PartnerFormatError('В заголовке партии IEK должна быть полная дата поступления.')
                            incoming_columns.append((i, datetime.strptime(match[1], '%d.%m.%Y').date(), label))
                    else:
                        for label in ('Артикул поставщика', 'Свободный остаток', 'Остаток', 'Зарезервировано'): col(label)
                        incoming_columns = [(col('СЭ в пути 24.09'), date(2026, 9, 24), 'СЭ в пути 24.09')]
                for n, row in enumerate(rows, header_row + 1):
                    if all(v is None for v in row): continue
                    seen += 1
                    sku = row[sku_col]
                    if sku is None: continue  # Explicit report subheaders / blank separator rows.
                    if isinstance(sku, str) and sku.strip().casefold() in ('итого', 'всего'): continue
                    if not isinstance(sku, str) or not sku.strip() or sku.startswith('#'):
                        rejected.append(n)
                        continue
                    p = item(sku)
                    if any(isinstance(v, str) and v.startswith(('#N/A', '#VALUE!', '#REF!', '#DIV/0!')) for v in row):
                        p['warnings'].add('В исходном Excel есть ошибка формулы; значение оставлено неизвестным.')
                    src = origin(n)
                    if kind == 'moq':
                        rule = number(row[rule_col])
                        if 'catalog' in p:
                            if p['rule'] == rule and p['article'] == str(row[article_col] or ''):
                                p['warnings'].add('Повтор MOQ с тем же кодом, артикулом и правилом учтён один раз.')
                                continue
                            p['rule'] = None
                            p['warnings'].add('Противоречащие строки MOQ для одного кода; правило неизвестно, заказ заблокирован.')
                            continue
                        p.update(catalog=True, name=str(row[name_col] or ''), article=str(row[article_col] or ''), rule=rule, rule_source=src)
                        p['warnings'].add('Кратность и минимальный заказ приняты равными значению из MOQ; проверьте перед утверждением.')
                    elif kind in ('monthly_stock', 'monthly_sales'):
                        p.setdefault('name', str(row[col('Номенклатура')] or ''))
                        if kind == 'monthly_stock':
                            unit = row[col('Ед.' if supplier == 'iek' else 'Ед.изм')]
                            if unit: p['units'].add(str(unit))
                        for idx, month in wide:
                            qty = number(row[idx])
                            if kind == 'monthly_stock':
                                p['stocks'].append({'month': month.isoformat(), 'quantity': qty, 'source': src})
                            else:
                                if month in p['monthly']:
                                    raise PartnerFormatError('В месячных продажах повторяется код 1С. Устраните дубликат.')
                                if qty is not None and qty < 0:
                                    p['warnings'].add('Отрицательная месячная продажа требует сверки возвратов; спрос за этот месяц неизвестен.')
                                    qty = None
                                p['monthly'][month] = MonthlySale(month=month, quantity=qty, source=src)
                    elif kind == 'incoming':
                        if 'incoming_seen' in p:
                            if p['incoming_raw'] == row:
                                p['warnings'].add('Одинаковая строка товара в пути учтена один раз.')
                            else:
                                p['incoming_complete'] = False
                                p.pop('inventory', None)
                                p['warnings'].add('Противоречащие строки товара в пути; требуется сверка перед заказом.')
                            continue
                        p['incoming_seen'] = True
                        p['incoming_raw'] = row
                        p.setdefault('name', str(row[col('Наименование')] or ''))
                        p.setdefault('article', str(row[col('Артикул ИЭК' if supplier == 'iek' else 'Артикул поставщика')] or ''))
                        p['incoming_complete'] = True
                        for idx, eta, shipment_id in incoming_columns:
                            qty = number(row[idx])
                            if qty is None or qty < 0:
                                p['incoming_complete'] = False
                            elif qty > 0:
                                p['incoming'].append((shipment_id, qty, eta, src))
                        if supplier == 'systeme_electric':
                            p['inventory'] = Inventory(as_of=SNAPSHOT, available=number(row[col('Свободный остаток')]),
                                on_hand=number(row[col('Остаток')]), reserved=number(row[col('Зарезервировано')]), source=Source(origin=src))
                            p['category'] = str(row[col('Категория 2026')] or '')
                    elif kind == 'sales':
                        dt = row[col('Дата')]
                        dt = dt.date() if isinstance(dt, datetime) else datetime.strptime(str(dt), '%d.%m.%Y %H:%M:%S').date()
                        qty = number(row[col('Количество')])
                        if qty is None:
                            p['warnings'].add('В динамике есть пустое количество: детальность месяца не подтверждена.')
                            p['unknown_document'] = True
                            continue
                        # This layout contains signed warehouse movements, not unsigned demand.
                        document = str(row[col('Документ')] or '')
                        if not document.startswith(('Расходная накладная', 'Возврат')):
                            p['warnings'].add('Неизвестный вид документа в динамике: детальность месяца не подтверждена.')
                            p['unknown_document'] = True
                            continue
                        if row[col('Ед.')]: p['units'].add(str(row[col('Ед.')]))
                        p.setdefault('name', str(row[col('Номенклатура')] or ''))
                        if dt <= SNAPSHOT:
                            p['sales'].append(Sale(date=dt, document_id=str(row[col('Номер')]), quantity=-qty, source=src))
                    used += 1
            audit.append({'name': f.original_name, 'sha256': f.sha256, 'kind': kind, 'rows_seen': seen, 'rows_imported': used, 'rejected_identifier_rows': rejected})
        except PartnerFormatError:
            raise
        except (ValueError, TypeError, IndexError, KeyError) as exc:
            raise PartnerFormatError(f'Не удалось разобрать отчёт {kind}. Проверьте структуру, даты и числовые значения.') from exc
        finally:
            book.close()

    if 'inventory' in files:
        f = files['inventory']
        with f.path.open(encoding='utf-8-sig', newline='') as stream:
            reader = csv.DictReader(stream, delimiter=';')
            if reader.fieldnames != ['sku_1c', 'stock_unit', 'available_stock', 'as_of', 'warehouse_id']:
                raise PartnerFormatError('inventory.csv не соответствует шаблону столбцов.')
            seen = set()
            for row in reader:
                sku = row['sku_1c']
                if sku in seen or sku not in products or row['as_of'] != SNAPSHOT.isoformat() or row['warehouse_id'] != SCOPE:
                    raise PartnerFormatError('Проверьте код, дату, охват складов и отсутствие дубликатов в inventory.csv.')
                seen.add(sku)
                p = products[sku]
                if not row['stock_unit'] or (p['units'] and p['units'] != {row['stock_unit']}):
                    raise PartnerFormatError('Единицы актуального остатка отличаются от исходных файлов.')
                p['units'].add(row['stock_unit'])
                qty = number(row['available_stock'])
                if qty is None: raise PartnerFormatError('Актуальный свободный остаток не может быть пустым.')
                p['inventory'] = Inventory(as_of=SNAPSHOT, available=qty, source=Source(origin=f'inventory.csv:CSV:{reader.line_num}'))
        audit.append({'name': f.original_name, 'sha256': f.sha256, 'kind': 'inventory', 'rows_imported': len(seen)})

    normalized, stock_audit = [], {}
    for sku, p in sorted(products.items()):
        warnings = p['warnings']
        warnings.add('Нет идентификаторов клиентов: концентрация разовых заказов по покупателю не проверена.')
        warnings.add('Месячные остатки не определяют точные дни отсутствия товара. Упущенный спрос требует дневных данных.')
        warnings.add('Охват: все склады выгрузки. Помесячные итоги не распределены по отдельным складам.')
        unit = next(iter(p['units'])) if len(p['units']) == 1 else 'неизвестно'
        if unit == 'неизвестно': warnings.add('Единица товара отсутствует или расходится между источниками; заказ заблокирован.')
        order_unit = unit if unit == 'шт' else 'уточнить единицу заказа'
        if unit != 'шт': warnings.add('Для метража и упаковок нужен подтверждённый коэффициент пересчёта; заказ заблокирован.')
        sums = defaultdict(float)
        for s in p['sales']: sums[s.date.replace(day=1)] += max(0, s.quantity)
        complete = tuple(m for m, value in p['monthly'].items() if value.quantity is not None and m < SNAPSHOT.replace(day=1)
                         and abs(sums[m] - value.quantity) < 1e-6 and not p.get('unknown_document'))
        rule = p.get('rule')
        if rule is not None and rule <= 0: rule = None
        normalized.append(ProductInput(supplier_id=request.supplier_id, sku=sku, warehouse_id=SCOPE,
            name=p.get('name', ''), supplier_article=p.get('article'), stock_unit=unit, order_unit=order_unit,
            stock_units_per_order_unit=1.0 if unit == 'шт' else None,
            units_source=Source(origin='hackalem-2026-v1: единицы заказа совпадают со складскими для шт', status='assumed') if unit == 'шт' else None,
            moq=rule, order_multiple=rule, order_rules_source=Source(origin=p.get('rule_source', 'MOQ отсутствует'), status='assumed'),
            category_code=p.get('category'), policy=SupplierPolicy(lead_time_days=14, review_days=7, source=Source(origin='MVP: параметры расчёта', status='assumed')),
            sales=tuple(p['sales']), monthly_sales=tuple(p['monthly'].values()), detail_months=complete,
            inventory=p.get('inventory'), incoming=tuple(Shipment(shipment_id=id, quantity=q, eta=eta, unit=unit, source=Source(origin=src, status='assumed')) for id,q,eta,src in p['incoming']),
            incoming_complete=p.get('incoming_complete', False), seasonality=seasonality, warnings=tuple(sorted(warnings))))
        stock_audit[sku] = p['stocks']
    issues = [
        {'id': 'scope', 'title': 'Все склады', 'detail': 'Отдельный склад не выделяется: месячные отчёты предоставлены в общем охвате.', 'severity': 'info', 'count': len(normalized)},
        {'id': 'stock', 'title': 'Нет актуального свободного остатка', 'detail': 'Заказ заблокирован. Добавьте inventory.csv по шаблону; месячный начальный остаток не заменяет свободный.', 'severity': 'warning', 'count': sum(p.inventory is None or (p.inventory.available is None and (p.inventory.on_hand is None or p.inventory.reserved is None)) for p in normalized)},
        {'id': 'units', 'title': 'Уточнить единицы заказа', 'detail': 'Для метража и упаковок нужен подтверждённый пересчёт.', 'severity': 'warning', 'count': sum(p.stock_units_per_order_unit is None for p in normalized)},
    ]
    issues.extend({'id': f['kind'] + '-identifiers', 'title': 'Строки без текстового кода 1С',
        'detail': f['name'] + ': строки ' + ', '.join(map(str, f['rejected_identifier_rows'])) + '. Не сопоставлены с товаром; проверьте разделители и коды.',
        'severity': 'warning', 'count': len(f['rejected_identifier_rows'])} for f in audit if f.get('rejected_identifier_rows'))
    version = PROFILE + ':' + sha256(''.join(sorted(f.sha256 for f in files.values())).encode()).hexdigest()[:20]
    ds = make_dataset(normalized, name=f'{request.supplier_id} — выгрузки 1С от 22.09.2026', mode='real', as_of=SNAPSHOT, version=version,
        quality={'source_count': len(audit), 'mapped_skus': len(normalized), 'files': audit, 'issues': issues})
    return ds.model_copy(update={'data': {**ds.data, 'monthly_stock_audit': stock_audit}})
