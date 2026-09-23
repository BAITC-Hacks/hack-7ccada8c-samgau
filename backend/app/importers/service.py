"""Shared importer entry point with explicit, validated workbook manifests."""
from __future__ import annotations

import json
from datetime import date
from hashlib import sha256
from pathlib import PurePosixPath
from typing import Literal
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from pydantic import Field, model_validator

from ..contracts import DatasetPayload, ImportRequest
from ..engine.models import CategoryPolicy, Identifier, Model, Stockout, SupplierPolicy
from ..engine.service import make_dataset
from .assemble import assemble_products
from .xlsx import IEKAdapter, ImportBatch, SystemeElectricAdapter, WorkbookMapping

MAPPING_VERSION = 'qor-explicit-xlsx-v1'
METADATA_SHEET = '_QOR_IMPORT'
SOURCE_KINDS = frozenset({'sales', 'monthly_sales', 'monthly_stock', 'incoming', 'seasonality', 'moq'})
MAX_ARCHIVE_BYTES = 512 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 5000


def canonical_supplier(value: str) -> str:
    aliases = {'iek': 'iek', 'иэк': 'iek', 'systeme': 'systeme_electric',
               'systeme_electric': 'systeme_electric', 'systeme electric': 'systeme_electric'}
    try:
        return aliases[value.strip().casefold()]
    except KeyError:
        raise ValueError('unsupported supplier; expected IEK or Systeme Electric') from None


class Coverage(Model):
    sku: Identifier
    detail_months: tuple[date, ...] = ()
    incoming_complete: bool = False
    stockouts: tuple[Stockout, ...] = ()

    @model_validator(mode='after')
    def month_starts(self):
        if any(m.day != 1 for m in self.detail_months):
            raise ValueError('coverage months must start on day 1')
        return self


class ImportConfiguration(Model):
    as_of: date
    warehouse_id: Identifier
    policy: SupplierPolicy
    category_policies: dict[str, CategoryPolicy] = Field(default_factory=dict)
    coverage: tuple[Coverage, ...] = ()

    @model_validator(mode='after')
    def unique_coverage(self):
        ids = [c.sku for c in self.coverage]
        if len(ids) != len(set(ids)):
            raise ValueError('duplicate SKU in coverage')
        return self


class WorkbookManifest(Model):
    format: Literal['qor-explicit-xlsx-v1'] = MAPPING_VERSION
    supplier_id: Identifier
    mapping: WorkbookMapping
    # Dataset metadata occurs once, in the dedicated MOQ workbook.
    configuration: ImportConfiguration | None = None


def write_manifest(workbook, manifest: WorkbookManifest):
    """Called by the synthetic generator or an explicit mapping preparation tool."""
    if METADATA_SHEET in workbook.sheetnames:
        raise ValueError('manifest already exists')
    encoded = manifest.model_dump_json()
    if len(encoded) > 32767:
        raise ValueError('manifest exceeds Excel cell capacity')
    sheet = workbook.create_sheet(METADATA_SHEET)
    sheet.append([MAPPING_VERSION])
    sheet.append([encoded])
    sheet.sheet_state = 'hidden'


def _check_archive(path):
    """Bound decompression before openpyxl loads XML or shared strings."""
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise ValueError('XLSX archive has too many entries')
            if sum(entry.file_size for entry in entries) > MAX_ARCHIVE_BYTES:
                raise ValueError('XLSX expanded archive exceeds size limit')
            if any(entry.file_size > 1024 * 1024 and entry.file_size / max(1, entry.compress_size) > 1000 for entry in entries):
                raise ValueError('XLSX compression ratio exceeds limit')
            if any(entry.flag_bits & 1 for entry in entries):
                raise ValueError('encrypted XLSX entries are unsupported')
            if len({entry.filename for entry in entries}) != len(entries):
                raise ValueError('duplicate archive entries')
    except BadZipFile as error:
        raise ValueError('file is not a valid XLSX archive') from error


def _manifest(path) -> WorkbookManifest:
    _check_archive(path)
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if METADATA_SHEET not in book.sheetnames:
            raise ValueError('explicit workbook manifest missing; real supplier layout is not yet configured')
        sheet = book[METADATA_SHEET]
        if sheet['A1'].value != MAPPING_VERSION:
            raise ValueError('unknown workbook manifest version')
        encoded = sheet['A2'].value
        if not isinstance(encoded, str) or len(encoded) > 32767:
            raise ValueError('invalid workbook manifest JSON')
        manifest = WorkbookManifest.model_validate_json(encoded)
        if manifest.mapping.sheet == METADATA_SHEET:
            raise ValueError('data mapping cannot target metadata sheet')
        return manifest
    finally:
        book.close()


def import_dataset(request: ImportRequest) -> DatasetPayload:
    request = ImportRequest.model_validate(request)
    if request.mapping_version != MAPPING_VERSION:
        raise ValueError(f'unsupported mapping_version; expected {MAPPING_VERSION}')
    if not 1 <= len(request.files) <= 12:
        raise ValueError('expected 1..12 input files')
    supplier = canonical_supplier(request.supplier_id)
    batch, manifests, seen_hashes = ImportBatch(), {}, set()
    imported_files, duplicates = [], []
    for uploaded in sorted(request.files, key=lambda f: (f.original_name, f.sha256)):
        original_name = PurePosixPath(uploaded.original_name.replace('\\', '/')).name
        if not original_name.lower().endswith('.xlsx') or uploaded.path.suffix.lower() != '.xlsx':
            raise ValueError('this mapping supports XLSX only; CSV profile is not configured')
        digest = sha256()
        with uploaded.path.open('rb') as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b''):
                digest.update(chunk)
        if digest.hexdigest() != uploaded.sha256:
            raise ValueError('file hash does not match ImportRequest')
        if uploaded.sha256 in seen_hashes:
            duplicates.append(original_name)
            continue
        seen_hashes.add(uploaded.sha256)
        manifest = _manifest(uploaded.path)
        if canonical_supplier(manifest.supplier_id) != supplier:
            raise ValueError('workbook supplier does not match requested supplier')
        kind = manifest.mapping.kind
        if kind in manifests:
            raise ValueError(f'multiple competing canonical files for {kind}')
        if manifest.configuration is not None and kind != 'moq':
            raise ValueError('configuration is only allowed in the MOQ workbook')
        manifests[kind] = manifest
        adapter = IEKAdapter() if supplier == 'iek' else SystemeElectricAdapter()
        result = adapter.read(uploaded.path, manifest.mapping)
        # The API assigns temporary random names. Persist the actual original filename instead.
        records = tuple(r.model_copy(update={'source_ref': tuple(t.model_copy(update={'file': original_name})
                                                                 for t in r.source_ref)}) for r in result.records)
        for record in records:
            if record.kind != 'seasonality' and record.values.get('warehouse_id') != request.warehouse_id:
                raise ValueError('workbook warehouse does not match requested scope')
        batch.add(result.model_copy(update={'records': records}))
        imported_files.append({'name': original_name, 'sha256': uploaded.sha256, 'kind': kind,
                               'mapping_version': manifest.mapping.version,
                               'rows_seen': result.rows_seen, 'rows_imported': result.rows_imported,
                               'rows_skipped': result.rows_skipped,
                               'issues': [i.model_dump(mode='json') for i in result.issues]})
    missing = SOURCE_KINDS - manifests.keys()
    if missing:
        raise ValueError(f'missing required source types: {", ".join(sorted(missing))}')
    config = manifests['moq'].configuration
    if config is None:
        raise ValueError('MOQ workbook requires explicit import configuration')
    if config.as_of != request.as_of or config.warehouse_id != request.warehouse_id:
        raise ValueError('import configuration date/warehouse differs from request')
    coverage = {(supplier, c.sku, config.warehouse_id): c for c in config.coverage}
    products = assemble_products(batch, as_of=request.as_of, policies={supplier: config.policy},
        category_policies={(supplier, code): policy for code, policy in config.category_policies.items()},
        detail_coverage={key: c.detail_months for key, c in coverage.items()},
        incoming_coverage={key: c.incoming_complete for key, c in coverage.items()},
        stockout_intervals={key: c.stockouts for key, c in coverage.items()})
    actual_skus = {p.sku for p in products}
    if {c.sku for c in config.coverage} - actual_skus:
        raise ValueError('coverage refers to unknown products')
    # Preserve the exact requested external supplier ID (e.g. IEK) required by captain validation.
    products = tuple(p.model_copy(update={'supplier_id': request.supplier_id}) for p in products)
    quality = {
        'source_count': len(imported_files), 'mapped_skus': len(products),
        'sources': sorted(SOURCE_KINDS), 'files': imported_files,
        'duplicate_files_ignored': duplicates,
        'warnings': sorted({w for p in products for w in p.warnings}),
        'issues': [{'id': f'{f["kind"]}:{i}', 'title': issue['code'],
                    'detail': issue['message'], 'severity': 'warning', 'count': 1}
                   for f in imported_files for i, issue in enumerate(f['issues'])],
    }
    fingerprint = sha256(json.dumps({'files': sorted(seen_hashes), 'supplier': request.supplier_id,
        'warehouse': request.warehouse_id, 'as_of': request.as_of.isoformat(),
        'mapping_version': MAPPING_VERSION}, sort_keys=True).encode()).hexdigest()[:20]
    dataset = make_dataset(products, name=f'{request.supplier_id} — {request.warehouse_id}', mode='real',
                           as_of=request.as_of, version=f'{MAPPING_VERSION}:{fingerprint}', quality=quality)
    # Persist audit records within opaque data, not the lightweight public quality response.
    return dataset.model_copy(update={'data': {**dataset.data,
        'source_records': [r.model_dump(mode='json') for r in batch.records]}})
