from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from app.engine import calculate
from app.engine.demo import AS_OF, make_demo_products
from app.importers import IEKAdapter, ImportBatch, SystemeElectricAdapter, WorkbookMapping, assemble_products
from app.importers.demo_files import create_demo_workbooks, import_demo_workbooks
from app.importers.xlsx import WideColumn


@pytest.fixture(scope='module')
def demo(tmp_path_factory):
    mappings = create_demo_workbooks(tmp_path_factory.mktemp('synthetic'))
    batch, products = import_demo_workbooks(mappings)
    return mappings, batch, products


def workbook(tmp_path, headers, rows):
    path = tmp_path / 'synthetic.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.title = 'Data'
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    book.save(path)
    book.close()
    return path


def sale_mapping(**kwargs):
    return WorkbookMapping(kind='sales', version='test', sheet='Data',
        columns={'sku': 'Код', 'date': 'Дата', 'document_id': 'Документ', 'quantity': 'Количество'},
        defaults={'unit': 'шт', 'warehouse_id': 'synthetic-almaty'}, **kwargs)


def test_all_twelve_sources_import_and_match_engine(demo):
    mappings, batch, products = demo
    assert len(mappings) == len(batch.results) == 12
    assert all(result.rows_skipped == 0 for result in batch.results)
    assert batch.missing_sources('iek') == batch.missing_sources('systeme_electric') == []
    expected = {p.sku: calculate(p, as_of=AS_OF) for p in make_demo_products()}
    for product in products:
        actual = calculate(product, as_of=AS_OF)
        assert actual.recommended_qty == expected[product.sku].recommended_qty
        if actual.forecast_qty is None:
            assert expected[product.sku].forecast_qty is None
        else:
            assert actual.forecast_qty == pytest.approx(expected[product.sku].forecast_qty)
    assert len(products) == 8


def test_repeat_file_is_idempotent(demo):
    _, original, _ = demo
    batch = ImportBatch()
    result = original.results[0]
    batch.add(result)
    count = len(batch.records)
    duplicate = batch.add(result)
    assert duplicate.duplicate_file
    assert not duplicate.records
    assert len(batch.records) == count


def test_leading_zero_signed_quantity_empty_and_source_trace(tmp_path):
    path = workbook(tmp_path, ['Код', 'Дата', 'Документ', 'Количество'],
                    [['00123', '22.09.2026', 'A', '-2'], ['00456', '22.09.2026', 'B', None],
                     [None, None, None, 100], ['Итого', None, None, 100]])
    result = IEKAdapter().read(path, sale_mapping())
    assert result.rows_imported == 1
    assert result.rows_skipped == 3
    assert result.records[0].values['sku'] == '00123'
    assert result.records[0].values['quantity'] == -2
    assert any(i.code == 'negative_movement' for i in result.issues)
    trace = next(t for t in result.records[0].source_ref if t.field == 'quantity')
    assert (trace.file, trace.sheet, trace.row, trace.raw, trace.value) == ('synthetic.xlsx', 'Data', 2, '-2', -2)
    assert len(trace.sha256) == 64


def test_numeric_sku_uses_excel_zero_mask(tmp_path):
    path = workbook(tmp_path, ['Код', 'Дата', 'Документ', 'Количество'], [[123, '22.09.2026', 'A', 1]])
    book = load_workbook(path)
    book['Data']['A2'].number_format = '00000'
    book.save(path)
    book.close()
    result = IEKAdapter().read(path, sale_mapping())
    assert result.records[0].values['sku'] == '00123'


def test_formula_without_cache_is_not_executed_or_zero(tmp_path):
    path = workbook(tmp_path, ['Код', 'Дата', 'Документ', 'Количество'], [['001', '22.09.2026', 'A', '=1+1']])
    result = IEKAdapter().read(path, sale_mapping())
    assert result.rows_skipped == 1
    assert not result.records
    assert any(i.code == 'formula_without_cached_value' for i in result.issues)


def test_short_date_requires_explicit_year(tmp_path):
    path = workbook(tmp_path, ['Код', 'Дата', 'Документ', 'Количество'], [['001', '24.09', 'A', 1]])
    assert IEKAdapter().read(path, sale_mapping()).rows_skipped == 1
    result = IEKAdapter().read(path, sale_mapping(short_date_year=2026))
    assert result.records[0].values['date'] == '2026-09-24'


def test_unknown_monthly_stock_stays_unknown_and_does_not_infer_intervals(tmp_path):
    path = workbook(tmp_path, ['sku', '2026-07', '2026-08'], [['001', None, 0]])
    mapping = WorkbookMapping(kind='monthly_stock', version='test', sheet='Data', columns={'sku': 'sku'},
        defaults={'warehouse_id': 'synthetic-almaty', 'unit': 'шт'}, snapshot_semantics='month_start',
        wide=(WideColumn(column='2026-07', date=date(2026, 7, 1)), WideColumn(column='2026-08', date=date(2026, 8, 1))))
    result = IEKAdapter().read(path, mapping)
    assert [r.values['quantity'] for r in result.records] == [None, 0]
    batch = ImportBatch()
    batch.add(result)
    policy = make_demo_products()[0].policy
    product, = assemble_products(batch, as_of=AS_OF, policies={'iek': policy})
    assert product.stockouts == ()
    assert product.inventory is None
    assert any('monthly_zero_stock_not_exact_stockout' in w for w in product.warnings)


def test_incoming_wide_headers_have_explicit_dates(tmp_path):
    path = workbook(tmp_path, ['sku', 'В пути 24.09', 'В пути 01.10'], [['001', 50, None]])
    mapping = WorkbookMapping(kind='incoming', version='test', sheet='Data', columns={'sku': 'sku'},
        defaults={'warehouse_id': 'synthetic-almaty', 'unit': 'шт'},
        wide=(WideColumn(column='В пути 24.09', date=date(2026, 9, 24), shipment_id='a'),
              WideColumn(column='В пути 01.10', date=date(2026, 10, 1), shipment_id='b')))
    result = SystemeElectricAdapter().read(path, mapping)
    assert result.records[0].values['eta'] == '2026-09-24'
    assert result.records[1].values['quantity'] is None


def test_ambiguous_headers_fail_early(tmp_path):
    path = workbook(tmp_path, ['Код', 'Код', 'Дата', 'Документ', 'Количество'], [['001', '002', '22.09.2026', 'A', 1]])
    with pytest.raises(ValueError, match='2 matches'):
        IEKAdapter().read(path, sale_mapping())


def test_invalid_pack_size_becomes_missing(tmp_path):
    path = workbook(tmp_path, ['sku', 'step'], [['001', 0]])
    mapping = WorkbookMapping(kind='moq', version='test', sheet='Data', columns={'sku': 'sku', 'order_multiple': 'step'})
    result = IEKAdapter().read(path, mapping)
    assert result.records[0].values['order_multiple'] is None
    assert any(i.code == 'invalid_order_rule' for i in result.issues)


def test_unknown_warehouse_cannot_be_joined(tmp_path):
    path = workbook(tmp_path, ['sku', 'step'], [['001', 12]])
    mapping = WorkbookMapping(kind='moq', version='test', sheet='Data', columns={'sku': 'sku', 'order_multiple': 'step'})
    batch = ImportBatch()
    batch.add(IEKAdapter().read(path, mapping))
    with pytest.raises(ValueError, match='warehouse scope'):
        assemble_products(batch, as_of=AS_OF, policies={})


def test_rejected_sales_row_invalidates_claim_of_complete_coverage(tmp_path):
    path = workbook(tmp_path, ['Код', 'Дата', 'Документ', 'Количество'],
                    [['001', '10.08.2026', 'A', 10], ['001', '11.08.2026', 'B', None]])
    batch = ImportBatch()
    batch.add(IEKAdapter().read(path, sale_mapping()))
    product, = assemble_products(batch, as_of=AS_OF, policies={'iek': make_demo_products()[0].policy},
        detail_coverage={('iek', '001', 'synthetic-almaty'): (date(2026, 8, 1),)})
    assert product.detail_months == ()
    result = calculate(product, as_of=AS_OF)
    assert result.forecast_qty is None
    assert 'detail_coverage_revoked_due_to_rejected_rows' in result.warnings
