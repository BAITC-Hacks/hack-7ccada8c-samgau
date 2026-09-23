"""Synthetic layout fixtures only; original company files never enter the repository."""
import csv
from datetime import date
from hashlib import sha256
import io
import time

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest

from app.config import Settings
from app.contracts import ImportFile, ImportRequest, CalculationParameters
from app.engine.service import calculate
from app.importers.partner import PROFILE, SCOPE, PartnerFormatError, import_partner
from app.main import create_app


@pytest.fixture
def source_files(tmp_path):
    rows = {
        'MOQ SystemElectric.xlsx': [['№','Номенклатура','Номенклатура.Код','Артикул','Кратность'],[1,'Тестовый товар','0001_','TEST-1',6]],
        'Ежемесячные остатки SystemElectric.xlsx': [['№','Номенклатура','Номенклатура.Код','Ед.изм','июнь 2026','июль 2026','авг. 2026'],[1,'Тестовый товар','0001_','шт',0,10,15]],
        'Ежемесячные продажи SystemElectric.xlsx': [['Номенклатура','Номенклатура.Код','июнь 2026','июль 2026','авг. 2026'],['Тестовый товар','0001_',90,100,110]],
        'Товар в пути_SystemElectric.xlsx': [[None],['Код 1с','Наименование','Артикул поставщика','Свободный остаток','Остаток','Зарезервировано','СЭ в пути 24.09','Категория 2026'],['0001_','Тестовый товар','TEST-1',0,0,0,0,'1']],
        'Динамика продаж_SystemElectric.xlsx': [['Дата','Номер','Документ','Код','Номенклатура','Ед.','Склад','Количество'],['01.08.2026 12:00:00','doc1','Расходная накладная 1','0001_','Тестовый товар','шт','Алматы',-110]],
        'Сезонность SystemElectric.xlsx': [[None] for _ in range(9)] + [[None]*11+['СЕЗОННОСТЬ']] + [[None]*11+[1.0] for _ in range(12)],
    }
    result=[]
    for name, data in rows.items():
        workbook=Workbook()
        if name.startswith('Товар'): workbook.active.title='TDSheet'
        for row in data: workbook.active.append(row)
        path=tmp_path/name
        workbook.save(path)
        result.append(ImportFile(path=path,original_name=name,sha256=sha256(path.read_bytes()).hexdigest()))
    return result


def request(files, **kwargs):
    return ImportRequest(files=files,supplier_id='systeme_electric',mapping_version=PROFILE,as_of=date(2026,9,22),warehouse_id=SCOPE,**kwargs)


def test_original_layouts_match_codes_and_keep_signed_sales(source_files):
    ds=import_partner(request(source_files))
    p=ds.data['products'][0]
    assert p['sku']=='0001_'
    assert p['sales'][0]['quantity']==110
    assert p['inventory']['available']==0  # observed zero survives
    assert p['incoming_complete'] is True
    assert p['detail_months']==['2026-08-01']
    result=calculate(ds,CalculationParameters(supplier_id='systeme_electric',as_of=ds.as_of))
    assert result.recommendations[0].recommended_qty>0
    assert result.recommendations[0].recommended_qty%6==0
    assert not result.recommendations[0].approval_blockers


def test_scope_and_missing_source_are_not_silently_inferred(source_files):
    req=request(source_files)
    with pytest.raises(PartnerFormatError): import_partner(req.model_copy(update={'warehouse_id':'Алматы'}))
    with pytest.raises(PartnerFormatError): import_partner(req.model_copy(update={'files':source_files[:-1]}))
    with pytest.raises(PartnerFormatError): import_partner(req.model_copy(update={'as_of':date(2026,9,23)}))


@pytest.mark.parametrize('cell,value,blocker', [
    ('D3', None, 'current_available_stock_missing'),
    ('G3', None, 'incoming_coverage_not_confirmed'),
    ('D3', 30, 'inventory_conflict_requires_confirmation'),
])
def test_missing_or_conflicting_inventory_cannot_be_approved(source_files, cell, value, blocker):
    original=next(f for f in source_files if f.original_name.startswith('Товар'))
    book=load_workbook(original.path)
    book.active[cell]=value
    if cell == 'D3' and value is None:
        book.active['E3']=None  # Neither explicit free stock nor on-hand is known.
    book.save(original.path)
    modified=original.model_copy(update={'sha256':sha256(original.path.read_bytes()).hexdigest()})
    files=[modified if f.path==original.path else f for f in source_files]
    ds=import_partner(request(files))
    row=calculate(ds,CalculationParameters(supplier_id='systeme_electric',as_of=ds.as_of)).recommendations[0]
    assert blocker in row.approval_blockers
    assert row.data_status=='blocked'


def test_confirmed_inventory_csv_changes_the_calculation(source_files,tmp_path):
    baseline=import_partner(request(source_files))
    path=tmp_path/'inventory.csv'
    path.write_text('sku_1c;stock_unit;available_stock;as_of;warehouse_id\n0001_;шт;500;2026-09-22;Все склады\n',encoding='utf-8-sig')
    extra=ImportFile(path=path,original_name=path.name,sha256=sha256(path.read_bytes()).hexdigest())
    changed=import_partner(request([*source_files,extra]))
    params=CalculationParameters(supplier_id='systeme_electric',as_of=baseline.as_of)
    before=calculate(baseline,params).recommendations[0]
    after=calculate(changed,params).recommendations[0]
    assert before.recommended_qty>0
    assert after.recommended_qty==0
    assert after.available_stock==500


def test_upload_calculation_approval_and_csv_roundtrip(source_files,tmp_path):
    settings=Settings(data_dir=tmp_path/'server',admin_token='test-import-secret',ai_provider='disabled')
    with TestClient(create_app(settings)) as c:
        headers={'Authorization':'Bearer '+c.post('/api/sessions').json()['token']}
        payload={'supplier':'systeme_electric','mapping_version':PROFILE,'as_of':'2026-09-22','warehouse_id':SCOPE}
        files=[('files',(f.original_name,f.path.read_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')) for f in source_files]
        assert c.post('/api/imports',headers=headers,data=payload,files=files).status_code==403
        r=c.post('/api/imports',headers={**headers,'X-Admin-Token':'test-import-secret'},data=payload,files=files)
        assert r.status_code==202,r.text
        job=r.json()
        for _ in range(100):
            job=c.get('/api/imports/'+job['import_id'],headers=headers).json()
            if job['status'] not in ('queued','running'): break
            time.sleep(.02)
        assert job['status']=='completed',job
        other={'Authorization':'Bearer '+c.post('/api/sessions').json()['token']}
        assert c.get('/api/datasets/'+job['dataset_id']+'/quality',headers=other).status_code==404
        run=c.post('/api/runs',headers=headers,json={'dataset_id':job['dataset_id'],'supplier_id':'systeme_electric','as_of':'2026-09-22'}).json()
        draft=c.post('/api/orders',headers=headers,json={'run_id':run['run_id'],'supplier_id':'systeme_electric','skus':['0001_']}).json()
        path='/api/orders/'+draft['draft_id']
        assert c.get(path+'/export.csv',headers=headers).status_code==409
        assert c.post(path+'/approve',headers=headers,json={'version':draft['version']}).status_code==422
        ok=c.post(path+'/approve',headers=headers,json={'version':draft['version'],'acknowledge_warnings':True})
        assert ok.status_code==200,ok.text
        exported=c.get(path+'/export.csv',headers=headers)
        rows=list(csv.DictReader(io.StringIO(exported.content.decode('utf-8-sig')),delimiter=';'))
        assert rows[0]['sku_1c']=='0001_'
        assert rows[0]['stock_unit']=='шт'
        assert float(rows[0]['quantity'])%6==0
        assert float(rows[0]['stock_equivalent'])==float(rows[0]['quantity'])
        assert rows[0]['data_mode']=='real'
